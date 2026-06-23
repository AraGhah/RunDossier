#!/usr/bin/env python3
"""
load_pilot_contacts.py
Load / sync warm-DM pilot contacts from a CSV or XLSX tracker into the
Preframe `public.pilot_requests` table via the idempotent
`upsert_pilot_contact(jsonb)` RPC.

Design goals (production):
  * Idempotent  - re-running never creates duplicates (dedupe on lower(email),
                  enforced both client-side and by the DB unique index / RPC).
  * Safe        - blank cells never overwrite existing DB values (RPC COALESCE).
  * Validated   - required fields, email format, and enum values checked before
                  any network call; bad rows are reported, never silently dropped.
  * Resilient   - exponential backoff retry on transient (network / 5xx / 429)
                  failures; 4xx (except 429) fail fast.
  * Auditable   - structured logging; the service-role key is NEVER logged.
  * Testable    - all parsing/validation is pure and import-safe (no network,
                  no hard dependency on `requests` at import time).

Usage:
  python load_pilot_contacts.py --file ../data/pilot_contacts_template.csv --dry-run
  python load_pilot_contacts.py --file ../Preframe_Pilot_Funnel_Tracker.xlsx --sheet Funnel
Env (see .env.example):
  SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY, [PILOT_OWNER]
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import re
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

log = logging.getLogger("pilot_loader")

# ---------------------------------------------------------------------------
# Schema knowledge (mirrors the DB CHECK constraints)
# ---------------------------------------------------------------------------
ALLOWED_STATUS = {"new", "contacted", "in_progress", "delivered", "closed"}
ALLOWED_CHANNEL = {"linkedin", "email", "both"}
ALLOWED_SENTIMENT = {"positive", "neutral", "negative", "no_reply"}
ALLOWED_OUTCOME = {
    "pending", "no_response", "replied", "call_booked", "pilot_accepted",
    "pilot_delivered", "case_study", "referral", "declined",
}
DEFAULT_STATUS = "new"

# Canonical payload fields accepted by the RPC.
TEXT_FIELDS = [
    "name", "email", "company", "sequencer", "profile_url", "status",
    "relationship", "persona", "channel", "message_variant", "owner",
    "mutual", "reply_sentiment", "outcome", "notes",
]
DATE_FIELDS = [
    "contacted_at", "first_reply_at", "follow_up_1_at",
    "follow_up_2_at", "pilot_sent_at",
]
INT_FIELDS = ["warmth"]
ALL_FIELDS = TEXT_FIELDS + DATE_FIELDS + INT_FIELDS

# Map friendly tracker headers -> canonical field keys (case-insensitive).
HEADER_ALIASES = {
    "name": "name",
    "persona": "persona",
    "company": "company",
    "email": "email",
    "linkedin url": "profile_url",
    "linkedin": "profile_url",
    "profile url": "profile_url",
    "profile_url": "profile_url",
    "mutual / how warm": "mutual",
    "mutual": "mutual",
    "relationship": "relationship",
    "warmth (1-5)": "warmth",
    "warmth": "warmth",
    "channel": "channel",
    "variant": "message_variant",
    "message variant": "message_variant",
    "message_variant": "message_variant",
    "status": "status",
    "contacted (date)": "contacted_at",
    "contacted": "contacted_at",
    "contacted_at": "contacted_at",
    "follow-up 1 (date)": "follow_up_1_at",
    "follow up 1": "follow_up_1_at",
    "follow_up_1_at": "follow_up_1_at",
    "follow-up 2 (date)": "follow_up_2_at",
    "follow up 2": "follow_up_2_at",
    "follow_up_2_at": "follow_up_2_at",
    "replied (date)": "first_reply_at",
    "replied?": "first_reply_at",
    "replied": "first_reply_at",
    "first_reply_at": "first_reply_at",
    "reply sentiment": "reply_sentiment",
    "reply_sentiment": "reply_sentiment",
    "pilot sent (date)": "pilot_sent_at",
    "pilot sent": "pilot_sent_at",
    "pilot_sent_at": "pilot_sent_at",
    "outcome": "outcome",
    "owner": "owner",
    "notes": "notes",
    "sequencer": "sequencer",
}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ValidationError(ValueError):
    """Raised when a row cannot be turned into a valid payload."""


# ---------------------------------------------------------------------------
# Pure helpers (no I/O, no network) - covered by unit tests
# ---------------------------------------------------------------------------
def clean_str(v: Any) -> str | None:
    """Trim to a non-empty string, or None. Handles pandas/openpyxl NaN."""
    if v is None:
        return None
    if isinstance(v, float) and v != v:  # NaN
        return None
    s = str(v).strip()
    return s or None


def normalize_email(v: Any) -> str | None:
    s = clean_str(v)
    return s.lower() if s else None


def is_valid_email(email: str) -> bool:
    return bool(EMAIL_RE.match(email))


def normalize_header(h: Any) -> str | None:
    s = clean_str(h)
    if not s:
        return None
    key = s.lower().strip()
    return HEADER_ALIASES.get(key, key if key in ALL_FIELDS else None)


def coerce_warmth(v: Any) -> int | None:
    s = clean_str(v)
    if s is None:
        return None
    try:
        n = int(float(s))
    except (TypeError, ValueError):
        raise ValidationError(f"warmth must be an integer 1-5, got {v!r}")
    if not 1 <= n <= 5:
        raise ValidationError(f"warmth must be between 1 and 5, got {n}")
    return n


def coerce_date(v: Any, field_name: str) -> str | None:
    """Return an ISO string Postgres can cast to timestamptz, or None."""
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    s = clean_str(v)
    if s is None:
        return None
    candidates = ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S",
                  "%Y-%m-%dT%H:%M:%S%z", "%m/%d/%Y", "%m/%d/%y")
    for fmt in candidates:
        try:
            return datetime.strptime(s, fmt).isoformat()
        except ValueError:
            continue
    raise ValidationError(f"{field_name} is not a recognizable date: {v!r}")


def _check_enum(value: str | None, allowed: set[str], field_name: str) -> str | None:
    if value is None:
        return None
    low = value.lower()
    if low not in allowed:
        raise ValidationError(
            f"{field_name}={value!r} invalid; allowed: {sorted(allowed)}"
        )
    return low


def row_to_payload(row: dict[str, Any], default_owner: str | None = None) -> dict[str, Any]:
    """Convert one canonical-keyed row dict into a validated RPC payload.

    Raises ValidationError on bad data. Returns {} for a fully blank row
    (caller should skip silently).
    """
    payload: dict[str, Any] = {}

    name = clean_str(row.get("name"))
    email = normalize_email(row.get("email"))

    # A row with no person identity (no name AND no email) is unstarted
    # scaffolding (e.g. a prefilled persona slot) -> skip silently.
    if not name and not email:
        return {}

    if not name:
        raise ValidationError("name is required")
    if not email:
        raise ValidationError(f"email is required (name={name!r})")
    if not is_valid_email(email):
        raise ValidationError(f"email is not valid: {email!r}")

    payload["name"] = name
    payload["email"] = email

    for f in ("company", "sequencer", "profile_url", "relationship",
              "persona", "message_variant", "mutual", "notes"):
        val = clean_str(row.get(f))
        if val is not None:
            payload[f] = val

    payload["status"] = _check_enum(clean_str(row.get("status")), ALLOWED_STATUS, "status") or DEFAULT_STATUS
    ch = _check_enum(clean_str(row.get("channel")), ALLOWED_CHANNEL, "channel")
    if ch:
        payload["channel"] = ch
    sent = _check_enum(clean_str(row.get("reply_sentiment")), ALLOWED_SENTIMENT, "reply_sentiment")
    if sent:
        payload["reply_sentiment"] = sent
    out = _check_enum(clean_str(row.get("outcome")), ALLOWED_OUTCOME, "outcome")
    if out:
        payload["outcome"] = out

    warmth = coerce_warmth(row.get("warmth"))
    if warmth is not None:
        payload["warmth"] = warmth

    for f in DATE_FIELDS:
        d = coerce_date(row.get(f), f)
        if d is not None:
            payload[f] = d

    if "owner" not in payload:
        owner = clean_str(row.get("owner")) or default_owner
        if owner:
            payload["owner"] = owner

    return payload


def dedupe_key(payload: dict[str, Any]) -> str:
    return (payload.get("email") or "").lower()


def dedupe_payloads(payloads: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    """Collapse by email; later rows win (shallow-merge over earlier)."""
    merged: dict[str, dict[str, Any]] = {}
    dup = 0
    for p in payloads:
        k = dedupe_key(p)
        if k in merged:
            dup += 1
            merged[k].update(p)
        else:
            merged[k] = dict(p)
        merged[k]["email"] = k  # canonical lowercased email regardless of input casing
    return list(merged.values()), dup


# ---------------------------------------------------------------------------
# File reading
# ---------------------------------------------------------------------------
def _map_record(raw: dict[Any, Any]) -> dict[str, Any]:
    """Remap a raw header->value record onto canonical field keys."""
    out: dict[str, Any] = {}
    for header, value in raw.items():
        key = normalize_header(header)
        if key:
            out[key] = value
    return out


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [_map_record(r) for r in csv.DictReader(fh)]


def read_xlsx(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    from openpyxl import load_workbook  # local import: not needed for unit tests
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    headers = [h for h in header]
    records = []
    for r in rows:
        records.append(_map_record({headers[i]: r[i] for i in range(len(headers)) if i < len(r)}))
    wb.close()
    return records


def read_rows(path: str | Path, sheet: str | None = None) -> list[dict[str, Any]]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"input file not found: {p}")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx(p, sheet)
    if p.suffix.lower() in (".csv", ".tsv"):
        return read_csv(p)
    raise ValueError(f"unsupported file type: {p.suffix} (use .csv or .xlsx)")


# ---------------------------------------------------------------------------
# Network layer (isolated + injectable for tests)
# ---------------------------------------------------------------------------
@dataclass
class LoadResult:
    inserted_or_updated: int = 0
    failed: int = 0
    skipped_blank: int = 0
    invalid: int = 0
    duplicates_merged: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.failed == 0 and self.invalid == 0


def _default_transport(url: str, headers: dict[str, str], body: dict[str, Any], timeout: float):
    import requests  # local import keeps unit tests dependency-free
    resp = requests.post(url, headers=headers, data=json.dumps(body), timeout=timeout)
    return resp.status_code, resp.text, dict(resp.headers)


def upsert_contact(
    payload: dict[str, Any],
    *,
    base_url: str,
    service_key: str,
    transport: Callable[..., tuple[int, str, dict]] = _default_transport,
    max_retries: int = 4,
    base_backoff: float = 0.5,
    timeout: float = 20.0,
    sleep: Callable[[float], None] = time.sleep,
) -> tuple[bool, str]:
    """Call the upsert RPC for a single payload. Returns (ok, detail)."""
    url = f"{base_url.rstrip('/')}/rest/v1/rpc/upsert_pilot_contact"
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }
    body = {"p": payload}
    attempt = 0
    while True:
        attempt += 1
        try:
            status, text, resp_headers = transport(url, headers, body, timeout)
        except Exception as exc:  # noqa: BLE001 - transient network error
            if attempt > max_retries:
                return False, f"network error after {attempt} attempts: {exc}"
            _backoff(attempt, base_backoff, None, sleep)
            continue

        if 200 <= status < 300:
            return True, "ok"
        if status == 429 or 500 <= status < 600:
            if attempt > max_retries:
                return False, f"HTTP {status} after {attempt} attempts: {_trunc(text)}"
            retry_after = _parse_retry_after(resp_headers)
            _backoff(attempt, base_backoff, retry_after, sleep)
            continue
        # 4xx (validation/auth) - do not retry
        return False, f"HTTP {status}: {_trunc(text)}"


def _backoff(attempt: int, base: float, retry_after: float | None, sleep: Callable[[float], None]) -> None:
    delay = retry_after if retry_after is not None else base * (2 ** (attempt - 1))
    log.warning("transient failure; retry %d in %.1fs", attempt, delay)
    sleep(delay)


def _parse_retry_after(headers: dict[str, str] | None) -> float | None:
    if not headers:
        return None
    val = headers.get("Retry-After") or headers.get("retry-after")
    if not val:
        return None
    try:
        return float(val)
    except ValueError:
        return None


def _trunc(text: str, n: int = 300) -> str:
    text = (text or "").replace("\n", " ")
    return text if len(text) <= n else text[:n] + "..."


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------
def build_payloads(records: list[dict[str, Any]], default_owner: str | None) -> tuple[list[dict], LoadResult]:
    result = LoadResult()
    payloads: list[dict[str, Any]] = []
    for i, rec in enumerate(records, start=2):  # row 1 = header
        try:
            p = row_to_payload(rec, default_owner)
        except ValidationError as exc:
            result.invalid += 1
            msg = f"row {i}: {exc}"
            result.errors.append(msg)
            log.error(msg)
            continue
        if not p:
            result.skipped_blank += 1
            continue
        payloads.append(p)
    payloads, dup = dedupe_payloads(payloads)
    result.duplicates_merged = dup
    return payloads, result


def run(
    file: str,
    *,
    sheet: str | None,
    dry_run: bool,
    default_owner: str | None,
    base_url: str | None,
    service_key: str | None,
    limit: int | None = None,
    transport: Callable[..., tuple[int, str, dict]] = _default_transport,
) -> LoadResult:
    records = read_rows(file, sheet)
    log.info("read %d data row(s) from %s", len(records), file)
    payloads, result = build_payloads(records, default_owner)
    if limit:
        payloads = payloads[:limit]
    log.info(
        "prepared %d valid payload(s) | blank skipped=%d invalid=%d dupes merged=%d",
        len(payloads), result.skipped_blank, result.invalid, result.duplicates_merged,
    )

    if dry_run:
        for p in payloads:
            log.info("DRY-RUN would upsert: %s <%s> [%s]",
                     p.get("name"), p.get("email"), p.get("status"))
        return result

    if not base_url or not service_key:
        raise SystemExit("SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY are required for a live load")

    for p in payloads:
        ok, detail = upsert_contact(
            p, base_url=base_url, service_key=service_key, transport=transport
        )
        if ok:
            result.inserted_or_updated += 1
            log.info("upserted %s <%s>", p.get("name"), p.get("email"))
        else:
            result.failed += 1
            msg = f"{p.get('email')}: {detail}"
            result.errors.append(msg)
            log.error("upsert failed for %s", msg)
    return result


def _configure_logging(verbose: bool) -> None:
    logging.basicConfig(
        level=logging.DEBUG if verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
    )


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Load warm-DM pilot contacts into pilot_requests.")
    ap.add_argument("--file", required=True, help="CSV or XLSX path")
    ap.add_argument("--sheet", default="Funnel", help="XLSX sheet name (default: Funnel)")
    ap.add_argument("--dry-run", action="store_true", help="validate + preview, no writes")
    ap.add_argument("--owner", default=os.environ.get("PILOT_OWNER"), help="default owner")
    ap.add_argument("--limit", type=int, default=None, help="cap number of rows")
    ap.add_argument("--verbose", action="store_true")
    return ap.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    _configure_logging(args.verbose)
    result = run(
        args.file,
        sheet=args.sheet,
        dry_run=args.dry_run,
        default_owner=args.owner,
        base_url=os.environ.get("SUPABASE_URL"),
        service_key=os.environ.get("SUPABASE_SERVICE_ROLE_KEY"),
        limit=args.limit,
    )
    log.info(
        "SUMMARY: upserted=%d failed=%d invalid=%d blank=%d dupes=%d %s",
        result.inserted_or_updated, result.failed, result.invalid,
        result.skipped_blank, result.duplicates_merged,
        "(dry-run)" if args.dry_run else "",
    )
    if not result.ok:
        for e in result.errors:
            log.error("  - %s", e)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
